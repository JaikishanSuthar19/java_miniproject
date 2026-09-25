import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    DARK_NAVY = RGBColor(15, 23, 42)      # #0F172A
    DEEP_NAVY = RGBColor(2, 6, 23)        # #020617
    PRIMARY_BLUE = RGBColor(37, 99, 235)  # #2563EB
    LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
    BORDER_COLOR = RGBColor(226, 232, 240)# #E2E8F0
    TEXT_MAIN = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    ACCENT_GREEN = RGBColor(22, 163, 74)  # #16A34A
    WHITE = RGBColor(255, 255, 255)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="COLLEGE COURSE REGISTRATION SYSTEM"):
        # Header banner bar
        header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = WHITE
        header_bar.line.color.rgb = BORDER_COLOR

        tf = header_bar.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.1)
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_BLUE

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_NAVY)

    # Accent decorative box
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    accent.fill.solid()
    accent.fill.fore_color.rgb = DEEP_NAVY
    accent.line.color.rgb = PRIMARY_BLUE
    accent.line.width = Pt(1.5)

    # Title & Subtitle text box
    tb = s1.shapes.add_textbox(Inches(1.3), Inches(1.2), Inches(10.733), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "MINI PROJECT 11 • GROUP 11"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY_BLUE

    p1 = tf.add_paragraph()
    p1.text = "College Course Registration\nManagement System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_before = Pt(8)
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "A Modern SaaS-Inspired Desktop Application Built with Java Swing & Advanced Collections"
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(203, 213, 225)

    # Metadata Cards (Cohort, Batch, Subject, Team)
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(4.3), Inches(10.733), Inches(1.9))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = DARK_NAVY
    meta_box.line.color.rgb = RGBColor(51, 65, 85)

    mtf = meta_box.text_frame
    mtf.word_wrap = True
    mtf.margin_left = Inches(0.4)
    mtf.margin_top = Inches(0.2)

    mp1 = mtf.paragraphs[0]
    mp1.text = "ACADEMIC DETAILS & TEAM PROFILE"
    mp1.font.size = Pt(11)
    mp1.font.bold = True
    mp1.font.color.rgb = PRIMARY_BLUE

    mp2 = mtf.add_paragraph()
    mp2.text = "Subject: Java Programming     |     Cohort: Jeff Bezos     |     Batch: 2025-29"
    mp2.font.size = Pt(13)
    mp2.font.bold = True
    mp2.font.color.rgb = WHITE
    mp2.space_before = Pt(4)

    mp3 = mtf.add_paragraph()
    mp3.text = "Team Members: Jaikishan  •  Omkar  •  Sarthak  •  Danish"
    mp3.font.size = Pt(13)
    mp3.font.color.rgb = RGBColor(147, 197, 253)
    mp3.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: Problem Statement & Objectives
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "Problem Statement & Objectives", "PROJECT FOUNDATION")

    # Left Card: Problem Statement
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.4)
    tf1.margin_right = Inches(0.4)
    tf1.margin_top = Inches(0.4)

    p = tf1.paragraphs[0]
    p.text = "📌 Problem Statement"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    p = tf1.add_paragraph()
    p.text = "\nModern academic institutions require an automated, conflict-free, and transparent course enrollment system to eliminate manual paperwork and over-enrollment.\n\n" \
             "Key Challenges Addressed:\n" \
             "• Accidental course over-enrollment beyond classroom capacity.\n" \
             "• Duplicate course registrations by the same student.\n" \
             "• Disconnected administration of faculty and course allocations.\n" \
             "• Lack of immediate student schedule feedback and credit visibility."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED

    # Right Card: Academic Objectives
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.4)
    tf2.margin_right = Inches(0.4)
    tf2.margin_top = Inches(0.4)

    p = tf2.paragraphs[0]
    p.text = "🎯 Core Syllabus Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    items = [
        "1. Manage Student, Faculty, & Course information.",
        "2. Implement dynamic registration & cancellation workflows.",
        "3. Real-time course capacity monitoring & seat enforcement.",
        "4. Utilize Fixed Arrays (Course[]) for base course storage.",
        "5. Utilize LinkedList for chronological registration audit logs.",
        "6. Utilize HashMap for fast O(1) student/course lookups.",
        "7. Utilize TreeMap for self-sorted course catalog display.",
        "8. Build a dual-portal desktop GUI using Java Swing & FlatLaf."
    ]
    for item in items:
        p = tf2.add_paragraph()
        p.text = item
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: Data Structures Architecture (Crucial for Viva)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "Java Collections & Algorithmic Complexities", "DATA STRUCTURES MATRIX")

    ds_info = [
        ("Course[] (Array)", "Fixed Initial Storage", "O(1) Access", "Stores the predefined baseline college courses; resized dynamically with System.arraycopy when expanded.", PRIMARY_BLUE),
        ("LinkedList<Registration>", "Chronological Audit Log", "O(1) Append", "Maintains transaction logs in order of occurrence; ideal for tracking active and cancelled records.", ACCENT_GREEN),
        ("HashMap<String, T>", "O(1) Fast Lookups", "O(1) Search", "Provides instant constant-time lookup for Students, Courses, and Faculty by unique ID strings.", PRIMARY_BLUE),
        ("TreeMap<String, Course>", "Self-Sorting Catalog", "O(log N) Search", "Backed by a Red-Black Binary Search Tree; automatically keeps courses in alphabetical order by ID.", ACCENT_GREEN),
    ]

    for idx, (title, role, complexity, desc, color) in enumerate(ds_info):
        x = Inches(0.8 + idx * 2.98)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(2.78), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_NAVY

        p = tf.add_paragraph()
        p.text = role.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_before = Pt(4)

        # Complexity pill
        p = tf.add_paragraph()
        p.text = "⚡ Complexity: " + complexity
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(14)

        p = tf.add_paragraph()
        p.text = "\nWhy Used:\n" + desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(6)

    # ==========================================
    # SLIDE 4: Architecture & Package Structure
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Layered MVC Architecture & Tech Stack", "SYSTEM DESIGN")

    layers = [
        ("Presentation Layer (GUI)", "Java Swing & FlatLaf", "Dual dashboards (Student/Admin), 9 specialized sub-panels, custom TableCellRenderers & Editors, live tab auto-refresh."),
        ("Business Service Layer", "Service Interfaces & Logic", "CourseService, StudentService, FacultyService, RegistrationService enforcing capacity limits & credit caps."),
        ("Data Persistence Layer", "In-Memory DataStore", "Centralized repository pattern managing Arrays, LinkedLists, HashMaps, and TreeMaps without external DB lag."),
        ("Domain Entity Layer (POJO)", "Core Business Models", "Student, Course, Faculty, and Registration models with strictly encapsulated fields and mutators.")
    ]

    for idx, (layer, tech, desc) in enumerate(layers):
        y = Inches(1.6 + idx * 1.3)
        row = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.15))
        row.fill.solid()
        row.fill.fore_color.rgb = CARD_BG
        row.line.color.rgb = BORDER_COLOR

        tf = row.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        p.text = layer
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_BLUE

        p = tf.add_paragraph()
        p.text = "Tech: " + tech + "  |  " + desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(3)

    # ==========================================
    # SLIDE 5: Screenshot 1 - Split-Screen Modern Login
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Dual-Portal Authentication (LoginFrame)", "USER INTERFACE")

    # Left: Screenshot
    img_path = "screenshots/01_login_frame.png"
    if os.path.exists(img_path):
        s5.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    # Right: Description Card
    desc_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "✨ Login Highlights"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Modern SaaS Split-Screen: Dark Navy branding left panel with graduation icon.",
        "• Dual Role Routing: Automatically redirects to Admin Portal ('admin') or Student Dashboard.",
        "• Enter-Key Accessibility: JRootPane.setDefaultButton binds Enter key to trigger login instantly.",
        "• Demo Credentials: Displayed directly on form for effortless evaluator testing.",
        "• Graceful Error Handling: Modal alert dialogs on invalid username or password."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 6: Screenshot 2 - Student Dashboard Overview
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, LIGHT_BG)
    add_header(s6, "Student Dashboard & Live Academic Overview", "STUDENT EXPERIENCE")

    img_path = "screenshots/02_student_dashboard.png"
    if os.path.exists(img_path):
        s6.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "📊 Dashboard Features"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• 4 Live Statistic Cards: Total Courses, Available Courses, Registered Courses, and Registered Credits.",
        "• Recent Courses Audit: Displays up to 3 most recently enrolled courses with credits and faculty.",
        "• Active Tab Highlight: Sidebar tab turns royal blue (#2563EB) for clear navigation state.",
        "• Dynamic Profile Pill: Top-right avatar shows initial letter badge, name, department, and semester.",
        "• Auto-Refresh Sync: Counters recalculate dynamically whenever tab is opened."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 7: Screenshot 3 - Course Catalog & Capacity
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, LIGHT_BG)
    add_header(s7, "Available Courses & Capacity Enforcement", "COURSE DISCOVERY")

    img_path = "screenshots/03_student_available_courses.png"
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "🔍 Catalog & Status Pills"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Real-Time Seat Availability: Computes Available Seats = Capacity - Registered Students.",
        "• Status Pill CellRenderer: Green 'AVAILABLE' badge vs Red 'FULL' badge (e.g. CS104 40/40 filled).",
        "• Live Search Bar: Filters courses instantaneously across Course ID, Course Name, and Faculty.",
        "• Alphabetical Order: Automatically sorted by Course ID using TreeMap backing store.",
        "• Read-Only Table: Built with DefaultTableModel.isCellEditable() = false."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 8: Screenshot 4 - Course Registration & Credit Cap
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, LIGHT_BG)
    add_header(s8, "Course Registration & Academic Credit Cap", "ENROLLMENT RULES")

    img_path = "screenshots/04_student_register_course.png"
    if os.path.exists(img_path):
        s8.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "⚡ Enrollment Engine"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Dynamic Course Info Card: Selecting from JComboBox updates faculty, credits, and vacant seats.",
        "• Over-Enrollment Guard: 'Confirm Registration' button disables automatically if course is full.",
        "• Duplicate Prevention: Blocks repeated registration attempts with explanatory warning.",
        "• 18-Credit Maximum Cap: Enforces university semester credit limit to prevent student overload.",
        "• Immediate Seat Decrement: Enrolling automatically increments registered student count."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 9: Screenshot 5 - My Courses & Cancellation
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, LIGHT_BG)
    add_header(s9, "My Courses & Interactive Cancellation", "SCHEDULE MANAGEMENT")

    img_path = "screenshots/05_student_my_courses.png"
    if os.path.exists(img_path):
        s9.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "🗑️ Instant Seat Release"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Schedule Overview: Shows enrolled courses with credits and registration dates.",
        "• Live Credits Header: Real-time calculation of total enrolled credits.",
        "• Table Row ButtonEditor: Custom TableCellEditor allows clicking red 'Cancel' button in table.",
        "• Confirmation Guard: Prompts student before cancelling to prevent accidental deregistration.",
        "• Restores Vacancy: Cancelling instantly decrements course enrollment count, releasing the seat."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 10: Screenshot 6 - Admin Dashboard Overview
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, LIGHT_BG)
    add_header(s10, "Admin Dashboard & Institutional Metrics", "ADMINISTRATIVE PORTAL")

    img_path = "screenshots/06_admin_dashboard.png"
    if os.path.exists(img_path):
        s10.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "🏛️ Executive Overview"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• 5 Institutional Metrics: Total Students, Total Courses, Total Faculty, Total Registrations, and Open Seats.",
        "• Real-Time Seat Aggregation: Loops through all courses to compute vacant seat capacity.",
        "• Quick Action Buttons: One-click navigation to Manage Courses, View Students, Faculty, and Registrations.",
        "• Auto-Refresh Synchronized: RefreshData() recalibrates counts whenever any admin tab is opened.",
        "• Administrator Profile Badge: Highlights system administrator role in top header."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 11: Screenshot 7 - Admin Course CRUD & Safe Deletion
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, LIGHT_BG)
    add_header(s11, "Course CRUD & Safe Deletion Guard", "COURSE ADMINISTRATION")

    img_path = "screenshots/07_admin_manage_courses.png"
    if os.path.exists(img_path):
        s11.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "🛡️ CRUD & Data Safety"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Complete Course CRUD: Admins can create new courses, edit offerings, or remove courses.",
        "• Modal Dialog Forms: Add & Edit modals validate non-empty fields and positive numbers.",
        "• Capacity Constraint Rule: Cannot reduce course capacity below current enrolled student count.",
        "• Safe Deletion Guard: Blocks deletion if registeredStudents > 0, preventing orphaned records.",
        "• 3-Way Collection Sync: Updates Array, HashMap, and TreeMap simultaneously."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 12: Screenshot 9 - Faculty Management (Problem Statement Requirement)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, LIGHT_BG)
    add_header(s12, "Faculty Management System", "INSTRUCTOR DIRECTORY")

    img_path = "screenshots/09_admin_faculty.png"
    if os.path.exists(img_path):
        s12.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "👨‍🏫 Faculty Portal"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Fulfills Requirement: Specifically fulfills the Problem Statement directive to manage faculty.",
        "• Faculty Directory: Displays Faculty ID, Name, Department (CSE/IT), and Email address.",
        "• Search Bar: Instant multi-field filtering by ID, Name, Department, or Email.",
        "• '+ Add Faculty' Modal: Validates inputs and stores into DataStore.faculties HashMap.",
        "• Duplicate ID Check: Rejects duplicate Faculty IDs with explanatory warning dialog."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 13: Screenshot 10 - Registration Audit Trail
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, LIGHT_BG)
    add_header(s13, "System-Wide Audit Log & Transactions", "AUDIT & MONITORING")

    img_path = "screenshots/10_admin_registrations.png"
    if os.path.exists(img_path):
        s13.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2))

    desc_box = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.2))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = CARD_BG
    desc_box.line.color.rgb = BORDER_COLOR
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_right = Inches(0.3)
    dtf.margin_top = Inches(0.4)

    p = dtf.paragraphs[0]
    p.text = "📋 Transaction History"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY

    features = [
        "• Complete Transaction Log: Tracks Registration ID, Student ID, Student Name, Course, and Date.",
        "• Status Badges: Displays green 'REGISTERED' for active enrollments and red 'CANCELLED' for history.",
        "• Preserves History: Cancelled enrollments are never hard-deleted; status transitions preserve audit integrity.",
        "• Powered by LinkedList: Demonstrates chronological sequential traversal of LinkedList<Registration>.",
        "• Refresh Control: Action button instantly queries latest registrations."
    ]
    for feat in features:
        p = dtf.add_paragraph()
        p.text = feat
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 14: Team Contributions (Group 11)
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, LIGHT_BG)
    add_header(s14, "Team Contributions & Module Ownership", "TEAM PROFILE")

    team = [
        ("Jaikishan", "Authentication & Student Portal", [
            "Student.java & StudentService.java",
            "LoginFrame.java with Enter key support",
            "StudentDashboard & CardLayout",
            "StudentHomePanel metric cards & live sync",
            "Main: Authentication + Navigation Architecture"
        ], PRIMARY_BLUE),
        ("Omkar", "Course Management & Safe CRUD", [
            "Course.java with capacity rules",
            "CourseService.java & array resizing",
            "CoursePanel.java with status badges",
            "AdminCoursePanel.java with Add/Edit forms",
            "Main: Course Service + Safe Deletion Guard"
        ], ACCENT_GREEN),
        ("Sarthak", "Registration Engine & Credit Cap", [
            "Registration.java & RegistrationService.java",
            "18-Credit Maximum Cap rule enforcement",
            "RegistrationPanel.java with info card",
            "MyCoursesPanel.java with table button editor",
            "Main: Enrollment Engine + Interactive Table UI"
        ], PRIMARY_BLUE),
        ("Danish", "Admin Architecture & Faculty System", [
            "DataStore.java (Collections Architecture)",
            "Faculty.java & FacultyService.java",
            "AdminFacultyPanel.java directory & modal",
            "AdminDashboard & AdminHomePanel (5 stats)",
            "Main: DataStore Architecture + Admin & Faculty"
        ], ACCENT_GREEN),
    ]

    for idx, (name, role, tasks, color) in enumerate(team):
        x = Inches(0.8 + idx * 2.98)
        card = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(2.78), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = "👨‍💻 " + name
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = DARK_NAVY

        p = tf.add_paragraph()
        p.text = role.upper()
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = "\nKey Deliverables:"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        for task in tasks:
            p = tf.add_paragraph()
            p.text = "• " + task
            p.font.size = Pt(10.5)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(3)

    # ==========================================
    # SLIDE 15: Conclusion & Future Scope (Dark Theme)
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, DARK_NAVY)

    tb = s15.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "PROJECT SUMMARY & CONCLUSION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    p = tf.add_paragraph()
    p.text = "A Complete, Enterprise-Grade Academic Solution"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(8)
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• Successfully demonstrates Object-Oriented Programming, Data Structures, and modern GUI design principles.\n" \
             "• Accurately fulfills all 8 Objectives and 10 Steps specified in Mini Project 11 syllabus.\n" \
             "• Features enterprise safeguards: Safe Course Deletion Guard, 18-Credit Maximum Cap, and live cross-panel sync.\n\n" \
             "🔮 Future Scope:\n" \
             "• Persistent SQL Database Integration (MySQL / PostgreSQL via JDBC).\n" \
             "• Automatic Waitlist Processing when enrolled students cancel their seats.\n" \
             "• Automated Student PDF Schedule Generation with iText."
    p.font.size = Pt(14.5)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # Output file
    output_path = "College_Course_Registration_Presentation_Group11.pptx"
    prs.save(output_path)
    print("SUCCESSFULLY_GENERATED: " + output_path)

if __name__ == "__main__":
    create_deck()
